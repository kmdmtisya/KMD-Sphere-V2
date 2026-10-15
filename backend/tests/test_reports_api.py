"""/api/v1/reports (P14-T15) through the API: figures equal the screens', owner-scoped, audited."""

import io
import uuid
from collections.abc import AsyncIterator
from datetime import UTC, datetime, timedelta
from typing import Any

import asyncpg
import httpx
import pytest
from fastapi import FastAPI
from openpyxl import load_workbook

from app.core.config import Settings
from app.main import create_app
from tests import auth_helpers as ah
from tests.authz_harness import TwoUsers, two_users  # noqa: F401
from tests.db_fixtures import migrated_settings  # noqa: F401

pytestmark = [pytest.mark.integration, pytest.mark.security]
TODAY = datetime.now(UTC).date()


@pytest.fixture(scope="module")
async def app(migrated_settings: Settings) -> AsyncIterator[FastAPI]:  # noqa: F811
    application = create_app(migrated_settings, readiness_checks={}, token_verifier=ah.verifier())
    yield application
    await application.state.db.dispose()


@pytest.fixture
async def http(app: FastAPI) -> AsyncIterator[httpx.AsyncClient]:
    transport = httpx.ASGITransport(app=app, raise_app_exceptions=False)
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as c:
        yield c


@pytest.fixture(scope="module")
async def db(migrated_settings: Settings) -> AsyncIterator[asyncpg.Connection]:  # noqa: F811
    conn = await asyncpg.connect(migrated_settings.database_dsn, timeout=5)
    yield conn
    await conn.close()


async def funded_portfolio(http: httpx.AsyncClient, h: dict[str, str]) -> str:
    pid = (
        await http.post(
            "/api/v1/portfolios",
            headers=h,
            json={"name": f"Rep {uuid.uuid4().hex[:6]}", "base_currency": "USD"},
        )
    ).json()["id"]
    r = await http.post(
        f"/api/v1/portfolios/{pid}/transactions",
        headers=h,
        json={
            "transaction_type": "DEPOSIT",
            "gross_amount": {"amount": "2500.00", "currency": "USD"},
            "trade_date": (TODAY - timedelta(days=3)).isoformat(),
            "currency": "USD",
        },
    )
    assert r.status_code == 201, r.text
    return str(pid)


async def report(
    http: httpx.AsyncClient, h: dict[str, str], kind: str, **params: Any
) -> httpx.Response:
    return await http.get(f"/api/v1/reports/{kind}", headers=h, params=params)


async def test_the_catalogue_lists_the_reports_and_formats(
    http: httpx.AsyncClient,
    two_users: TwoUsers,  # noqa: F811
) -> None:
    assert (await http.get("/api/v1/reports")).status_code == 401
    items = (await http.get("/api/v1/reports", headers=two_users.alice.headers)).json()
    assert {i["kind"] for i in items} == {
        "portfolio",
        "performance",
        "net-worth",
        "income",
        "goals",
    }
    assert all(i["formats"] == ["csv", "xlsx", "pdf"] for i in items)
    assert {i["kind"] for i in items if i["needs_portfolio"]} == {"portfolio", "performance"}


async def test_a_portfolio_report_equals_the_summary_in_every_format(
    http: httpx.AsyncClient,
    two_users: TwoUsers,  # noqa: F811
) -> None:
    h = two_users.alice.headers
    pid = await funded_portfolio(http, h)
    summary = (await http.get(f"/api/v1/portfolios/{pid}/summary", headers=h)).json()

    csv = await report(http, h, "portfolio", portfolio_id=pid, format="csv")
    assert csv.status_code == 200 and csv.headers["content-type"].startswith("text/csv")
    assert (
        "attachment" in csv.headers["content-disposition"]
        and ".csv" in csv.headers["content-disposition"]
    )
    assert csv.headers["cache-control"] == "no-store"
    text = csv.content.decode("utf-8-sig")
    assert f"Total value,{summary['total_value']['amount']}" in text
    assert f"Cash,{summary['cash']['amount']}" in text
    assert "Currency,USD" in text and "not investment, tax or legal advice" in text

    xlsx = await report(http, h, "portfolio", portfolio_id=pid, format="xlsx")
    book = load_workbook(io.BytesIO(xlsx.content))
    totals = {r[0].value: r[1].value for r in book["Totals"].iter_rows(min_row=2)}
    assert totals["Total value"] == float(summary["total_value"]["amount"])

    pdf = await report(http, h, "portfolio", portfolio_id=pid, format="pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")


async def test_net_worth_income_goals_and_performance_reports_build(
    http: httpx.AsyncClient,
    two_users: TwoUsers,  # noqa: F811
) -> None:
    h = two_users.alice.headers
    pid = await funded_portfolio(http, h)
    goal = await http.post(
        "/api/v1/goals",
        headers=h,
        json={
            "name": "Report goal",
            "goal_type": "wealth",
            "target": {"amount": "10000.00", "currency": "USD"},
            "target_date": (TODAY + timedelta(days=900)).isoformat(),
            "target_in_todays_money": False,
            "assumptions": {
                "contribution": {"amount": "100.00", "currency": "USD"},
                "returns": {"base": "5"},
            },
        },
    )
    assert goal.status_code == 201, goal.text

    nw = (await http.get("/api/v1/net-worth", headers=h, params={"currency": "USD"})).json()
    r = await report(http, h, "net-worth", format="csv", currency="USD")
    assert r.status_code == 200, r.text
    assert f"Net worth,{nw['net_worth']['amount']}" in r.content.decode("utf-8-sig")

    perf = (await http.get(f"/api/v1/portfolios/{pid}/performance", headers=h)).json()
    r = await report(http, h, "performance", portfolio_id=pid, format="csv")
    assert f"Total value,{perf['total_value']['amount']}" in r.content.decode("utf-8-sig")

    r = await report(http, h, "goals", format="csv")
    assert "Report goal" in r.content.decode("utf-8-sig")

    income = (
        await http.get("/api/v1/income/summary", headers=h, params={"currency": "USD"})
    ).json()
    r = await report(http, h, "income", format="csv", currency="USD")
    assert f"Per year (net),{income['yearly']['amount']}" in r.content.decode("utf-8-sig")
    r = await report(http, h, "income", format="pdf", portfolio_id=pid)
    assert r.status_code == 200 and r.content.startswith(b"%PDF")


async def test_reports_are_the_users_own_and_validated(
    http: httpx.AsyncClient,
    two_users: TwoUsers,  # noqa: F811
) -> None:
    alice, bob = two_users.alice.headers, two_users.bob.headers
    pid = await funded_portfolio(http, alice)
    for kind in ("portfolio", "performance"):
        r = await report(http, bob, kind, portfolio_id=pid, format="csv")
        assert r.status_code == 404, (kind, r.status_code)
        assert "2500" not in r.text
    assert (await report(http, bob, "income", portfolio_id=pid, format="csv")).status_code == 404
    assert (
        await report(http, alice, "portfolio", format="csv")
    ).status_code == 422  # needs a portfolio
    assert (await report(http, alice, "nonsense", format="csv")).status_code == 404
    assert (await report(http, alice, "goals", format="docx")).status_code == 422
    assert (await report(http, alice, "goals", format="csv", currency="usd")).status_code == 422
    assert (await http.get(f"/api/v1/reports/portfolio?portfolio_id={pid}")).status_code == 401


async def test_each_generated_file_is_audited_without_figures(
    http: httpx.AsyncClient,
    two_users: TwoUsers,  # noqa: F811
    db: asyncpg.Connection,
) -> None:
    h = two_users.alice.headers
    pid = await funded_portfolio(http, h)
    assert (await report(http, h, "portfolio", portfolio_id=pid, format="csv")).status_code == 200
    rows = await db.fetch(
        "SELECT details::text AS d, resource_id FROM audit_events "
        "WHERE action = 'reports.generated' AND resource_id = $1",
        pid,
    )
    assert len(rows) == 1
    assert '"kind": "portfolio"' in rows[0]["d"] and '"format": "csv"' in rows[0]["d"]
    assert "2500" not in rows[0]["d"]
