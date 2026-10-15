"""Report files (P14-T15): the CSV, XLSX and PDF renderers of app/modules/reports/render.py."""

import io
from datetime import UTC, datetime

from openpyxl import load_workbook

from app.modules.reports.model import DISCLOSURES, Report, Table
from app.modules.reports.render import render, safe_text, to_csv, to_pdf, to_xlsx

NOW = datetime(2026, 10, 16, 9, 30, tzinfo=UTC)


def report(**over: object) -> Report:
    table = Table(
        "Holdings",
        ("Asset", "Quantity", "Value"),
        [("Acme Corp", "10.5", "1234.50"), ('=HYPERLINK("http://evil.test")', "1", "-5.00")],
        frozenset({1, 2}),
    )
    base = {
        "kind": "portfolio",
        "title": "Portfolio statement: Main",
        "generated_at": NOW,
        "currency": "USD",
        "as_of": NOW,
        "tables": [table],
        "notes": ["One position is unpriced."],
    }
    return Report(**(base | over))  # type: ignore[arg-type]


def test_text_that_a_spreadsheet_could_run_is_neutralised() -> None:
    for risky in ("=SUM(A1)", "+1", "-1+2", "@cmd", "\tx"):
        assert safe_text(risky).startswith("'")
    assert safe_text("Acme Corp") == "Acme Corp"


def test_csv_has_the_header_block_the_tables_and_the_disclosures() -> None:
    data = to_csv(report())
    assert data.startswith(b"\xef\xbb\xbf")  # byte-order mark for spreadsheets
    text = data.decode("utf-8-sig")
    assert "Report,Portfolio statement: Main" in text
    assert "Generated,2026-10-16 09:30 UTC" in text
    assert "Currency,USD" in text and "Data as of,2026-10-16 09:30 UTC" in text
    assert "Acme Corp,10.5,1234.50" in text
    for line in ("One position is unpriced.", *DISCLOSURES):
        assert line in text


def test_csv_neutralises_a_formula_in_text_but_leaves_negative_numbers_alone() -> None:
    text = to_csv(report()).decode("utf-8-sig")
    assert "'=HYPERLINK" in text  # the asset name is text
    assert ",1,-5.00" in text  # a negative amount in a numeric column is untouched


def test_xlsx_opens_with_one_sheet_per_table_and_numbers_as_numbers() -> None:
    book = load_workbook(io.BytesIO(to_xlsx(report())))
    assert book.sheetnames == ["Summary", "Holdings"]
    summary = {
        row[0].value: row[1].value for row in book["Summary"].iter_rows(min_row=1, max_row=4)
    }
    assert summary["Currency"] == "USD"
    holdings = book["Holdings"]
    assert [c.value for c in holdings[1]] == ["Asset", "Quantity", "Value"]
    assert holdings["B2"].value == 10.5 and holdings["C2"].value == 1234.5
    assert holdings["C3"].value == -5
    assert str(holdings["A3"].value).startswith("'=HYPERLINK")  # still text, never a formula
    assert holdings["A3"].data_type != "f"


def test_xlsx_sheet_names_are_valid_and_unique() -> None:
    awkward = [Table("A/B: [x]?", ("c",), [("1",)]), Table("A/B: [x]?", ("c",), [("2",)])]
    names = load_workbook(io.BytesIO(to_xlsx(report(tables=awkward)))).sheetnames
    assert len(set(names)) == 3 and all(not set("[]:*?/\\") & set(n) for n in names)


def test_pdf_is_a_pdf_and_user_text_cannot_use_the_markup() -> None:
    nasty = report(
        tables=[Table("T", ("Asset",), [("<b>bold</b> & <i>x",)])], title="A <script> title"
    )
    data = to_pdf(nasty)
    assert data.startswith(b"%PDF") and len(data) > 1000
    assert render(report(), "pdf").startswith(b"%PDF")


def test_an_empty_table_still_renders_in_every_format() -> None:
    empty = report(tables=[Table("Nothing", ("A", "B"), [])])
    for fmt in ("csv", "xlsx", "pdf"):
        assert len(render(empty, fmt)) > 100
