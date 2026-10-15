"""Renders a Report as CSV, XLSX or PDF (P14-T15, DEC-33). Pure: bytes in, bytes out.

CSV: UTF-8 with a byte-order mark (so spreadsheets read accents), the header block, then each table
under its title, separated by a blank line. Text that a spreadsheet could run as a formula (it
starts with = + - @ or a control character) is prefixed with an apostrophe, so an asset name cannot
inject a formula (CSV injection). Numeric columns are never altered.

XLSX: one sheet per table plus a Summary sheet with the header block and disclosures; numeric
columns are written as numbers (Decimal text, not rounded) with a bold header row.

PDF: the header block, each table, then the disclosures."""

import csv
import io
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, TableStyle
from reportlab.platypus import Table as PdfTable

from app.modules.reports.model import Report, Table

_FORMULA_START = ("=", "+", "-", "@", "\t", "\r", "\n")
CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def safe_text(value: str) -> str:
    """Text that a spreadsheet will show as text, never run as a formula."""
    return "'" + value if value.startswith(_FORMULA_START) else value


def _cell(table: Table, index: int, value: str) -> str:
    return value if index in table.numeric else safe_text(value)


def to_csv(report: Report) -> bytes:
    out = io.StringIO(newline="")
    writer = csv.writer(out, lineterminator="\r\n")
    for label, value in report.facts():
        writer.writerow([label, safe_text(value)])
    for table in report.tables:
        writer.writerow([])
        writer.writerow([safe_text(table.title)])
        writer.writerow([safe_text(c) for c in table.columns])
        for row in table.rows:
            writer.writerow([_cell(table, i, v) for i, v in enumerate(row)])
    writer.writerow([])
    writer.writerow(["Notes"])
    for line in report.disclosures():
        writer.writerow([safe_text(line)])
    return b"\xef\xbb\xbf" + out.getvalue().encode("utf-8")


def _number(value: str) -> Decimal | str:
    try:
        return Decimal(value)
    except InvalidOperation:
        return value  # an empty or non-numeric cell stays text


_SHEET_FORBIDDEN = str.maketrans(dict.fromkeys("[]:*?/\\", " "))


def to_xlsx(report: Report) -> bytes:
    book = Workbook()
    summary = book.active
    if summary is None:
        raise RuntimeError("a new workbook has an active sheet")
    summary.title = "Summary"
    bold = Font(bold=True)
    for label, value in report.facts():
        summary.append([label, safe_text(value)])
        summary.cell(row=summary.max_row, column=1).font = bold
    summary.append([])
    for line in report.disclosures():
        summary.append([safe_text(line)])
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 60
    used = {"Summary"}
    for table in report.tables:
        name = table.title.translate(_SHEET_FORBIDDEN)[:28] or "Table"
        base, n = name, 2
        while name in used:
            name, n = f"{base[:25]} {n}", n + 1
        used.add(name)
        sheet = book.create_sheet(name)
        sheet.append([safe_text(c) for c in table.columns])
        for cell in sheet[1]:
            cell.font = bold
        for row in table.rows:
            sheet.append(
                [_number(v) if i in table.numeric else safe_text(v) for i, v in enumerate(row)]
            )
        for i, column in enumerate(table.columns, start=1):
            sheet.column_dimensions[sheet.cell(row=1, column=i).column_letter].width = max(
                14, min(40, len(column) + 4)
            )
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def to_pdf(report: Report) -> bytes:
    styles = getSampleStyleSheet()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=report.title,
        author="WealthSphere",
    )
    story: list[object] = [Paragraph(_esc(report.title), styles["Title"])]
    for label, value in report.facts()[1:]:
        story.append(Paragraph(f"<b>{_esc(label)}:</b> {_esc(value)}", styles["Normal"]))
    for table in report.tables:
        story.append(Spacer(1, 6 * mm))
        story.append(Paragraph(_esc(table.title), styles["Heading2"]))
        if not table.rows:
            story.append(Paragraph("No entries.", styles["Normal"]))
            continue
        data = [[Paragraph(f"<b>{_esc(c)}</b>", styles["Normal"]) for c in table.columns]]
        for row in table.rows:
            data.append([Paragraph(_esc(v), styles["Normal"]) for v in row])
        grid = PdfTable(data, repeatRows=1)
        style = [
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]
        for i in table.numeric:
            style.append(("ALIGN", (i, 1), (i, -1), "RIGHT"))
        grid.setStyle(TableStyle(style))
        story.append(grid)
    story.append(Spacer(1, 8 * mm))
    story.append(Paragraph("Notes", styles["Heading2"]))
    for line in report.disclosures():
        story.append(Paragraph(_esc(line), styles["Normal"]))
    doc.build(story)
    return buffer.getvalue()


def _esc(text: str) -> str:
    """PDF paragraphs read a small markup: user text (asset names) must not be able to use it."""
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def render(report: Report, fmt: str) -> bytes:
    if fmt == "csv":
        return to_csv(report)
    if fmt == "xlsx":
        return to_xlsx(report)
    if fmt == "pdf":
        return to_pdf(report)
    raise ValueError(f"unknown format {fmt}")
